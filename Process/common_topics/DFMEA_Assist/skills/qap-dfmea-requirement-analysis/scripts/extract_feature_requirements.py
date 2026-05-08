# (c) Qorix 2026
"""Extract implemented requirements for a feature from TSYN SRS.
This tool reads the Qorix TSYN requirement workbook and returns the subset of
requirements that:
- are linked to the given Feature ID ("Feature Id" column contains the Feature ID), and
- have "Requirement Implementation Status" equal to "Implemented" or
  "Implemented Partially".
Before requirements are extracted, the tool verifies that the feature itself is
implemented by checking the "Features" sheet:
- The row with ID == <feature_id> must have "Feature Implementation Status" equal
  to "Implemented" or "Implemented Partially".
If the feature is not implemented (or not found), an empty JSON array is
returned.
Expected workbook:
- Path: Input/Qorix_AP_TSYN_SRS.xlsx (default)
- Sheets:
  - Features
  - Requirement Analysis
Output:
- JSON array of objects with keys:
  - Requirement Id
  - Requirement Description
  - Requirement Category
Dependencies:
- openpyxl
"""
from __future__ import annotations
import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple
try:
    from openpyxl import load_workbook
    from openpyxl.worksheet.worksheet import Worksheet
except ModuleNotFoundError as exc:  # pragma: no cover
    raise SystemExit(
        "Missing dependency 'openpyxl'. Install it with: py -m pip install openpyxl"
    ) from exc

_ALLOWED_IMPLEMENTATION_STATUSES = {
    "implemented",
    "implemented partially",
}

@dataclass(frozen=True)
class _SheetSpec:
    """Sheet configuration with required headers."""
    name: str
    required_headers: Tuple[str, ...]

_FEATURES_SHEET = _SheetSpec(
    name="Features",
    required_headers=(
        "ID",
        "Feature Implementation Status",
    ),
)
_REQUIREMENTS_SHEET = _SheetSpec(
    name="Requirement Analysis",
    required_headers=(
        "ID",
        "Requirement Description",
        "Requirement Category",
        "Requirement Implementation Status",
        "Feature Id",
    ),
)

def _to_text(value: Any) -> str:
    """Convert a worksheet cell value to trimmed text.
    Args:
        value: Any cell value coming from openpyxl (None, str, numbers, etc.).
    Returns:
        A trimmed string representation (empty string for None).
    """
    if value is None:
        return ""
    return str(value).strip()

def _normalize_header(header: str) -> str:
    """Normalize a header name for case/spacing-insensitive matching.
    Normalization strategy:
    - lower-case
    - collapse whitespace
    - remove all non-alphanumeric characters
    Args:
        header: Header cell text.
    Returns:
        Normalized header token.
    """
    header = " ".join(_to_text(header).lower().split())
    return re.sub(r"[^a-z0-9]", "", header)

def _normalize_status(status: str) -> str:
    """Normalize a status cell for matching against allowed values."""
    return " ".join(_to_text(status).lower().split())

def _is_implemented(status: str) -> bool:
    """Check whether a status indicates implementation (fully or partially)."""
    return _normalize_status(status) in _ALLOWED_IMPLEMENTATION_STATUSES

def _parse_feature_ids(cell_text: str) -> Sequence[str]:
    """Parse the 'Feature Id' cell into individual feature IDs.
    The TSYN SRS uses comma-separated feature IDs (e.g. "AP-8351, AP-8352"). This
    function splits on commas/semicolons/newlines and strips whitespace.
    Args:
        cell_text: The raw cell content.
    Returns:
        A list of feature IDs.
    """
    text = _to_text(cell_text)
    if not text or text == "--":
        return []
    tokens = [t.strip() for t in re.split(r"[,;\n]+", text) if t.strip()]
    return tokens

def _find_header_row(
    ws: Worksheet,
    required_headers: Sequence[str],
    *,
    max_scan_rows: int = 30,
) -> Tuple[int, Mapping[str, int]]:
    """Locate the header row and return a column index map.
    The workbook may contain a title row before the column headers. This function
    scans the first `max_scan_rows` rows and selects the first row that contains
    all required headers.
    Args:
        ws: Worksheet.
        required_headers: Exact header names expected to exist in the sheet.
        max_scan_rows: Maximum number of initial rows to scan.
    Returns:
        Tuple of (header_row_index_1based, header_to_col_index_0based).
    Raises:
        ValueError: If a suitable header row cannot be found.
    """
    required_norm = {_normalize_header(h) for h in required_headers}
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True), start=1
    ):
        if not row or not any(v is not None and _to_text(v) for v in row):
            continue
        norm_to_idx: Dict[str, int] = {}
        for col_idx, cell in enumerate(row):
            norm = _normalize_header(_to_text(cell))
            if norm:
                norm_to_idx[norm] = col_idx
        if required_norm.issubset(norm_to_idx.keys()):
            # Map original required header names to their 0-based column indices.
            header_map = {
                header: norm_to_idx[_normalize_header(header)] for header in required_headers
            }
            return row_idx, header_map
    raise ValueError(
        f"Header row not found in sheet '{ws.title}'. Required headers: {list(required_headers)}"
    )

def _is_empty_row(values: Sequence[Any]) -> bool:
    """Return True if all row values are empty/None after trimming."""
    return not any(v is not None and _to_text(v) for v in values)

def extract_feature_requirements(
    *,
    xlsx_path: Path,
    feature_id: str,
    features_sheet_name: str = _FEATURES_SHEET.name,
    requirements_sheet_name: str = _REQUIREMENTS_SHEET.name,
) -> List[Dict[str, str]]:
    """Extract implemented requirements linked to a feature.
    Args:
        xlsx_path: Path to the requirement workbook.
        feature_id: Feature ID to search for (e.g. "AP-8359").
        features_sheet_name: Sheet name containing feature list and status.
        requirements_sheet_name: Sheet name containing requirements and feature mapping.
    Returns:
        A JSON-serializable list of requirements. If the feature is not
        implemented (or not found), returns an empty list.
    Raises:
        FileNotFoundError: If the workbook does not exist.
        ValueError: If required sheets/headers are missing.
    """
    feature_id = _to_text(feature_id)
    if not feature_id:
        return []
    if not xlsx_path.exists():
        raise FileNotFoundError(str(xlsx_path))
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    if features_sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Missing sheet '{features_sheet_name}'. Available: {wb.sheetnames}"
        )
    if requirements_sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Missing sheet '{requirements_sheet_name}'. Available: {wb.sheetnames}"
        )
    features_ws = wb[features_sheet_name]
    features_header_row, features_cols = _find_header_row(
        features_ws, _FEATURES_SHEET.required_headers
    )
    # 1) Verify feature implementation status.
    feature_is_implemented = False
    for row in features_ws.iter_rows(
        min_row=features_header_row + 1,
        values_only=True,
    ):
        if not row or _is_empty_row(row):
            continue
        row_feature_id = _to_text(row[features_cols["ID"]])
        if row_feature_id != feature_id:
            continue
        feature_status = _to_text(row[features_cols["Feature Implementation Status"]])
        feature_is_implemented = _is_implemented(feature_status)
        break
    if not feature_is_implemented:
        return []
    # 2) Filter requirements for the feature.
    req_ws = wb[requirements_sheet_name]
    req_header_row, req_cols = _find_header_row(req_ws, _REQUIREMENTS_SHEET.required_headers)
    results: List[Dict[str, str]] = []
    for row in req_ws.iter_rows(min_row=req_header_row + 1, values_only=True):
        if not row or _is_empty_row(row):
            continue
        req_status = _to_text(row[req_cols["Requirement Implementation Status"]])
        if not _is_implemented(req_status):
            continue
        mapped_feature_ids = _parse_feature_ids(row[req_cols["Feature Id"]])
        if feature_id not in mapped_feature_ids:
            continue
        requirement_id = _to_text(row[req_cols["ID"]])
        requirement_description = _to_text(row[req_cols["Requirement Description"]])
        requirement_category = _to_text(row[req_cols["Requirement Category"]])
        results.append(
            {
                "Requirement Id": requirement_id,
                "Requirement Description": requirement_description,
                "Requirement Category": requirement_category,
            }
        )
    return results

def _parse_args(argv: Optional[Sequence[str]]) -> argparse.Namespace:
    """Parse CLI arguments."""
    parser = argparse.ArgumentParser(
        description=(
            "Extract requirements for a Feature ID from Input/Qorix_AP_TSYN_SRS.xlsx "
            "(TSYN functional cluster)."
        )
    )
    parser.add_argument(
        "feature_id",
        help='Feature ID to search for (e.g. "AP-8359").',
    )
    parser.add_argument(
        "--xlsx",
        default="Input/Qorix_AP_TSYN_SRS.xlsx",
        help="Path to the requirement workbook (.xlsx). Default: %(default)s",
    )
    parser.add_argument(
        "--pretty",
        action="store_true",
        help="Pretty-print JSON output (indent=2).",
    )
    return parser.parse_args(argv)

def main(argv: Optional[Sequence[str]] = None) -> int:
    """CLI entrypoint.
    Args:
        argv: Optional argument vector (without program name). If None, uses sys.argv[1:].
    Returns:
        Process exit code.
    """
    args = _parse_args(argv)
    try:
        requirements = extract_feature_requirements(
            xlsx_path=Path(args.xlsx),
            feature_id=args.feature_id,
        )
    except (FileNotFoundError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    json_kwargs = {"ensure_ascii": False}
    if args.pretty:
        json_kwargs.update({"indent": 2, "sort_keys": False})
    print(json.dumps(requirements, **json_kwargs))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
