# (c) Qorix 2026
"""List implemented (and partially implemented) TSYN features.
This tool reads the Qorix TSYN requirement workbook and returns the subset of
features that have "Feature Implementation Status" equal to:
- Implemented
- Implemented Partially
Expected workbook:
- Path: Input/Qorix_AP_TSYN_SRS.xlsx (default)
- Sheet:
  - Features
Output:
- JSON array of objects with keys:
  - Feature Id
  - Feature Description
  - Feature Implementation Status
Dependencies:
- openpyxl
"""
from __future__ import annotations
import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple
try:
    from openpyxl import load_workbook
    from openpyxl.worksheet.worksheet import Worksheet
except ModuleNotFoundError as exc:  # pragma: no cover
    raise SystemExit(
        "Missing dependency 'openpyxl'. Install it with: py -m pip install openpyxl"
    ) from exc

_ALLOWED_IMPLEMENTATION_STATUSES = {
    "implemented",
    "implemented partially",
}

@dataclass(frozen=True)
class _SheetSpec:
    """Sheet configuration with required headers."""
    name: str
    required_headers: Tuple[str, ...]

_FEATURES_SHEET = _SheetSpec(
    name="Features",
    required_headers=(
        "ID",
        "Feature Implementation Status",
    ),
)

def _to_text(value: Any) -> str:
    """Convert a worksheet cell value to trimmed text.
    Args:
        value: Any cell value coming from openpyxl (None, str, numbers, etc.).
    Returns:
        A trimmed string representation (empty string for None).
    """
    if value is None:
        return ""
    return str(value).strip()

def _normalize_header(header: str) -> str:
    """Normalize a header name for case/spacing-insensitive matching."""
    header = " ".join(_to_text(header).lower().split())
    return re.sub(r"[^a-z0-9]", "", header)

def _normalize_status(status: str) -> str:
    """Normalize a status cell for matching against allowed values."""
    return " ".join(_to_text(status).lower().split())

def _is_implemented(status: str) -> bool:
    """Check whether a status indicates implementation (fully or partially)."""
    return _normalize_status(status) in _ALLOWED_IMPLEMENTATION_STATUSES

def _find_header_row(
    ws: Worksheet,
    required_headers: Sequence[str],
    *,
    max_scan_rows: int = 30,
) -> Tuple[int, Mapping[str, int], Mapping[str, int]]:
    """Locate the header row and return a column index map.
    The workbook may contain a title row before the column headers. This function
    scans the first `max_scan_rows` rows and selects the first row that contains
    all required headers.
    Args:
        ws: Worksheet.
        required_headers: Exact header names expected to exist in the sheet.
        max_scan_rows: Maximum number of initial rows to scan.
    Returns:
        Tuple of:
        - header_row_index_1based
        - required header -> col index (0-based)
        - normalized header -> col index (0-based) for all headers in the row
    Raises:
        ValueError: If a suitable header row cannot be found.
    """
    required_norm = {_normalize_header(h) for h in required_headers}
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True), start=1
    ):
        if not row or not any(v is not None and _to_text(v) for v in row):
            continue
        norm_to_idx: Dict[str, int] = {}
        for col_idx, cell in enumerate(row):
            norm = _normalize_header(_to_text(cell))
            if norm:
                norm_to_idx[norm] = col_idx
        if required_norm.issubset(norm_to_idx.keys()):
            header_map = {
                header: norm_to_idx[_normalize_header(header)] for header in required_headers
            }
            return row_idx, header_map, norm_to_idx
    raise ValueError(
        f"Header row not found in sheet '{ws.title}'. Required headers: {list(required_headers)}"
    )

def _is_empty_row(values: Sequence[Any]) -> bool:
    """Return True if all row values are empty/None after trimming."""
    return not any(v is not None and _to_text(v) for v in values)

def list_implemented_features(
    *,
    xlsx_path: Path,
    features_sheet_name: str = _FEATURES_SHEET.name,
) -> List[Dict[str, str]]:
    """List features with status Implemented or Implemented Partially.
    Args:
        xlsx_path: Path to the TSYN requirement workbook.
        features_sheet_name: Sheet name containing feature list and status.
    Returns:
        JSON-serializable list of features.
    Raises:
        FileNotFoundError: If the workbook does not exist.
        ValueError: If required sheets/headers are missing.
    """
    if not xlsx_path.exists():
        raise FileNotFoundError(str(xlsx_path))
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    if features_sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Missing sheet '{features_sheet_name}'. Available: {wb.sheetnames}"
        )
    ws = wb[features_sheet_name]
    header_row, required_cols, norm_cols = _find_header_row(ws, _FEATURES_SHEET.required_headers)
    id_col = required_cols["ID"]
    status_col = required_cols["Feature Implementation Status"]
    desc_col = norm_cols.get(_normalize_header("Feature Description"))
    results: List[Dict[str, str]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or _is_empty_row(row):
            continue
        feature_id = _to_text(row[id_col])
        if not feature_id:
            continue
        status = _to_text(row[status_col])
        if not _is_implemented(status):
            continue
        description = ""
        if desc_col is not None and desc_col < len(row):
            description = _to_text(row[desc_col])
        results.append(
            {
                "Feature Id": feature_id,
                "Feature Description": description,
                "Feature Implementation Status": status,
            }
        )
    return results

def _parse_args(argv: Optional[Sequence[str]]) -> argparse.Namespace:
    """Parse CLI arguments."""
    parser = argparse.ArgumentParser(
        description=(
            "List features with status Implemented / Implemented Partially from "
            "Input/Qorix_AP_TSYN_SRS.xlsx (TSYN functional cluster)."
        )
    )
    parser.add_argument(
        "--xlsx",
        default="Input/Qorix_AP_TSYN_SRS.xlsx",
        help="Path to the requirement workbook (.xlsx). Default: %(default)s",
    )
    parser.add_argument(
        "--pretty",
        action="store_true",
        help="Pretty-print JSON output (indent=2).",
    )
    return parser.parse_args(argv)

def main(argv: Optional[Sequence[str]] = None) -> int:
    """CLI entrypoint.
    Args:
        argv: Optional argument vector (without program name). If None, uses sys.argv[1:].
    Returns:
        Process exit code.
    """
    args = _parse_args(argv)
    try:
        features = list_implemented_features(xlsx_path=Path(args.xlsx))
    except (FileNotFoundError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    json_kwargs = {"ensure_ascii": False}
    if args.pretty:
        json_kwargs.update({"indent": 2, "sort_keys": False})
    print(json.dumps(features, **json_kwargs))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
